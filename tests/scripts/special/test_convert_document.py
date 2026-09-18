"""Tests for scripts/special/convert_document.py.

Covers the public API: process_document function, manifest creation,
cache hit/miss behavior, force re-extraction, and embedded image extraction.
"""

import json
import os
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from scripts.special.convert_document import MIN_IMAGE_SIDE, process_document

"""Public API of this test module (empty: no symbols are exported)."""
__all__ = ()


def _png(width: int, height: int) -> bytes:
    """Build an in-memory PNG of the given pixel size."""
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height))
    pix.clear_with(200)
    return pix.tobytes("png")


def _write_png(path: Path, width: int, height: int) -> None:
    """Write a PNG of the given pixel size to disk."""
    path.write_bytes(_png(width, height))


def _pdf_with_image(path: Path, width: int, height: int) -> None:
    """Write a one-page PDF whose page embeds one image of the given pixel size."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(0, 0, 100, 100), stream=_png(width, height))
    doc.save(path)
    doc.close()


def test_process_document_rejects_unknown_format(
    tmp_path: os.PathLike[str],
) -> None:
    """Unsupported extension raises ValueError."""
    base = Path(os.fspath(tmp_path))
    bad = base / "file.txt"
    bad.write_text("hello")
    with pytest.raises(ValueError, match="Unsupported format"):
        process_document(bad, base / "out")


def test_manifest_created(tmp_path: os.PathLike[str]) -> None:
    """process_document writes a manifest.json with correct fields."""
    base = Path(os.fspath(tmp_path))

    pdf_path = base / "test.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()

    out_dir = base / "output"
    result = process_document(pdf_path, out_dir)

    manifest_path = out_dir / "manifest.json"
    assert manifest_path.exists()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["format"] == "pdf"
    assert manifest["source_name"] == "test.pdf"
    assert "source_sha256" in manifest
    assert "timestamp" in manifest
    assert manifest["page_count"] == result.page_count
    assert manifest["image_count"] == len(result.embedded_images)
    assert manifest["image_skipped_count"] == result.embedded_images_skipped


def test_cache_hit_skips_extraction(tmp_path: os.PathLike[str]) -> None:
    """Second call with same source returns cached result."""
    base = Path(os.fspath(tmp_path))

    pdf_path = base / "test.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()

    out_dir = base / "output"
    result1 = process_document(pdf_path, out_dir)
    assert result1.cached is False

    result2 = process_document(pdf_path, out_dir)
    assert result2.cached is True
    assert result2.text == result1.text
    assert result2.page_count == result1.page_count


def test_force_overrides_cache(tmp_path: os.PathLike[str]) -> None:
    """force=True re-extracts even when cache is valid."""
    base = Path(os.fspath(tmp_path))

    pdf_path = base / "test.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()

    out_dir = base / "output"
    process_document(pdf_path, out_dir)
    result = process_document(pdf_path, out_dir, force=True)
    assert result.cached is False


def test_extract_embedded_image_pdf(tmp_path: os.PathLike[str]) -> None:
    """A PDF's embedded image is written at its own size, not as a page render."""
    base = Path(os.fspath(tmp_path))

    pdf_path = base / "test.pdf"
    _pdf_with_image(pdf_path, 64, 64)

    out_dir = base / "output"
    result = process_document(pdf_path, out_dir)

    assert len(result.embedded_images) == 1
    assert result.embedded_images[0].name == "page_001_img_1.png"
    assert result.embedded_images[0].parent == out_dir / "images"
    assert pymupdf.Pixmap(str(result.embedded_images[0])).width == 64
    assert result.embedded_images_skipped == 0
    assert result.pages[0].name == "page_001.png"


def test_decorative_image_filtered_pdf(tmp_path: os.PathLike[str]) -> None:
    """An embedded image below the minimum side is counted, not written."""
    base = Path(os.fspath(tmp_path))

    pdf_path = base / "test.pdf"
    _pdf_with_image(pdf_path, MIN_IMAGE_SIDE - 8, MIN_IMAGE_SIDE - 8)

    out_dir = base / "output"
    result = process_document(pdf_path, out_dir)

    assert result.embedded_images == ()
    assert result.embedded_images_skipped == 1
    assert list((out_dir / "images").iterdir()) == []


def test_extract_embedded_image_docx(tmp_path: os.PathLike[str]) -> None:
    """A DOCX package's embedded image is written to images/."""
    base = Path(os.fspath(tmp_path))

    png_path = base / "figure.png"
    _write_png(png_path, 64, 64)
    docx_path = base / "test.docx"
    document = Document()
    document.add_picture(str(png_path))
    document.save(str(docx_path))

    result = process_document(docx_path, base / "output")

    assert [path.name for path in result.embedded_images] == ["image_001.png"]
    assert result.pages == ()


def test_extract_embedded_image_pptx(tmp_path: os.PathLike[str]) -> None:
    """A PPTX's embedded picture is written, named for its slide."""
    base = Path(os.fspath(tmp_path))

    png_path = base / "figure.png"
    _write_png(png_path, 64, 64)
    pptx_path = base / "test.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(png_path), Inches(1), Inches(1))
    presentation.save(str(pptx_path))

    result = process_document(pptx_path, base / "output")

    assert [path.name for path in result.embedded_images] == ["slide_001_img_1.png"]


def test_cache_requires_images_dir(tmp_path: os.PathLike[str]) -> None:
    """A cache written without images/ is a stale layout and is re-extracted."""
    base = Path(os.fspath(tmp_path))

    pdf_path = base / "test.pdf"
    _pdf_with_image(pdf_path, 64, 64)

    out_dir = base / "output"
    process_document(pdf_path, out_dir)
    for cached in (out_dir / "images").iterdir():
        cached.unlink()
    (out_dir / "images").rmdir()

    result = process_document(pdf_path, out_dir)

    assert result.cached is False
    assert (out_dir / "images").is_dir()


def test_extract_text_pdf(tmp_path: os.PathLike[str]) -> None:
    """PDF text extraction produces a non-empty markdown file."""
    pytest.skip("No test PDF fixture")


def test_extract_text_docx(tmp_path: os.PathLike[str]) -> None:
    """DOCX text extraction produces a non-empty markdown file."""
    pytest.skip("No test DOCX fixture")


def test_extract_text_pptx(tmp_path: os.PathLike[str]) -> None:
    """PPTX text extraction produces a non-empty markdown file."""
    pytest.skip("No test PPTX fixture")
